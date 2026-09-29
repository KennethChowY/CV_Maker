"""Downloads: a real PDF (printed by a Chrome-based browser already on the computer) and a Word file.

Both are made from the CV page as it currently looks, including hand edits,
hidden sections and the chosen template.
"""

from __future__ import annotations

import io
import os
import shutil
import subprocess
import sys
import tempfile
from html import escape
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

STATIC = Path(__file__).parent / "static"
PAGE_MM = {"A4": (210, 297), "letter": (215.9, 279.4)}


class ExportError(RuntimeError):
    pass


# ---- PDF ---------------------------------------------------------------

def find_browser() -> str | None:
    """A Chrome-based browser that can print to PDF without showing anything."""
    override = os.environ.get("CV_MAKER_BROWSER")
    if override:
        return override if Path(override).exists() or shutil.which(override) else None
    if sys.platform == "darwin":
        apps = ["Google Chrome.app/Contents/MacOS/Google Chrome", "Chromium.app/Contents/MacOS/Chromium",
                "Microsoft Edge.app/Contents/MacOS/Microsoft Edge", "Brave Browser.app/Contents/MacOS/Brave Browser",
                "Arc.app/Contents/MacOS/Arc"]
        for root in (Path("/Applications"), Path.home() / "Applications"):
            for app in apps:
                if (root / app).exists():
                    return str(root / app)
        return None
    if sys.platform.startswith("win"):
        roots = [os.environ.get(v) for v in ("PROGRAMFILES", "PROGRAMFILES(X86)", "LOCALAPPDATA")]
        rel = [r"Google\Chrome\Application\chrome.exe", r"Microsoft\Edge\Application\msedge.exe",
               r"BraveSoftware\Brave-Browser\Application\brave.exe"]
        for root in filter(None, roots):
            for r in rel:
                if Path(root, r).exists():
                    return str(Path(root, r))
        return None
    for name in ("google-chrome", "google-chrome-stable", "chromium", "chromium-browser", "microsoft-edge", "brave-browser"):
        found = shutil.which(name)
        if found:
            return found
    return None


PAGE_SLACK_MM = 8  # room left on each page for small print-vs-screen differences (matches app.js)


def page_document(body: str, *, title: str, template: str, accent: str, page_size: str, scale: float,
                  extra_css: str = "", paginate: bool = True) -> str:
    """A complete HTML page for printing: no page margin (so browsers add no header or footer),
    with the margin applied as padding repeated on every page instead."""
    css = (STATIC / "cv.css").read_text(encoding="utf-8")
    size = "letter" if page_size == "letter" else "A4"
    width_mm, height_mm = PAGE_MM[size]
    # Page headers and "continued" lines, planned with the same code as the preview.
    pages_js = (f"CVLayout.paginate(cv, CVLayout.mm({height_mm - 26 - PAGE_SLACK_MM}), "
                f"(cv.querySelector('.cv-name') || {{textContent: ''}}).textContent.trim());" if paginate else "")
    accent_css = f"--cv-accent: {accent}; " if template == "modern" else ""  # other templates are black
    layout_js = (STATIC / "cv-layout.js").read_text(encoding="utf-8")  # same line tightening as the preview
    return (
        f"<!doctype html><html><head><meta charset='utf-8'><title>{escape(title)}</title><style>{css}\n"
        f"@page {{ size: {size}; margin: 0; }}\nhtml, body {{ margin: 0; background: #fff; }}\n"
        f".cv {{ padding: 13mm 14mm; box-decoration-break: clone; -webkit-box-decoration-break: clone; "
        f"box-sizing: border-box; width: {width_mm}mm; }}\n"
        f"{extra_css}</style></head><body>"
        f"<article class='cv t-{template}' style='{accent_css}--cv-scale: {scale:.3f}'>{body}</article>"
        f"<script>{layout_js}\nconst cv = document.querySelector('.cv');\nCVLayout.tighten(cv);\n{pages_js}</script>"
        f"</body></html>"
    )


def html_to_pdf(document: str, browser: str | None = None) -> bytes:
    browser = browser or find_browser()
    if not browser:
        raise ExportError("No Chrome, Edge or Brave browser found to make the PDF.")
    with tempfile.TemporaryDirectory() as tmp:
        src, out = Path(tmp, "page.html"), Path(tmp, "page.pdf")
        src.write_text(document, encoding="utf-8")
        cmd = [browser, "--headless", "--disable-gpu", "--no-first-run", "--no-default-browser-check",
               "--no-pdf-header-footer", "--print-to-pdf-no-header", f"--user-data-dir={Path(tmp, 'profile')}",
               f"--print-to-pdf={out}", src.as_uri()]
        if sys.platform.startswith("linux") and hasattr(os, "geteuid") and os.geteuid() == 0:
            cmd.insert(1, "--no-sandbox")  # Chrome refuses to run sandboxed as root on Linux
        try:
            subprocess.run(cmd, capture_output=True, timeout=90, check=False)
        except (OSError, subprocess.TimeoutExpired) as e:
            raise ExportError(f"The browser couldn't make the PDF: {e}") from e
        if not out.exists() or out.stat().st_size == 0:
            raise ExportError("The browser couldn't make the PDF.")
        return out.read_bytes()


# ---- Word --------------------------------------------------------------

FONTS = {"classic": "Georgia", "modern": "Arial", "minimal": "Helvetica"}


def _text(node) -> str:
    return " ".join(node.get_text(" ", strip=True).split()) if node else ""


def _rule_below(paragraph, color: str = "999999") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", color)):
        bottom.set(qn(key), value)
    borders.append(bottom)
    p_pr.append(borders)


def _tight(paragraph, before: float = 0, after: float = 0) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    fmt.line_spacing = 1.08


def html_to_docx(body: str, *, template: str, accent: str, page_size: str) -> bytes:
    """A Word version of the CV page, following the same structure and template."""
    soup = BeautifulSoup(body, "html.parser")
    doc = Document()
    width_mm, height_mm = PAGE_MM["letter" if page_size == "letter" else "A4"]
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(width_mm), Mm(height_mm)
    section.left_margin = section.right_margin = Mm(16)
    section.top_margin = section.bottom_margin = Mm(14)
    text_width = Mm(width_mm - 32)

    normal = doc.styles["Normal"]
    normal.font.name = FONTS.get(template, "Arial")
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), normal.font.name)
    normal.font.size = Pt(10.5)
    ink = RGBColor(0x16, 0x19, 0x1D)
    muted = RGBColor(0x4F, 0x55, 0x61)
    heading_color = RGBColor.from_string(accent.lstrip("#").upper()) if template == "modern" else ink
    centred = template == "classic"

    def para(before=0, after=0, align=None):
        p = doc.add_paragraph()
        _tight(p, before, after)
        if align is not None:
            p.alignment = align
        return p

    def two_sides(left_runs, right_text, italic=False):
        p = para()
        p.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
        for text, bold in left_runs:
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
        if right_text:
            r = p.add_run("\t" + right_text)
            r.italic = italic
            r.font.color.rgb = muted
            r.font.size = Pt(9.5)
        return p

    header = soup.find(class_="cv-header")
    if header:
        p = para(align=WD_ALIGN_PARAGRAPH.CENTER if centred else None)
        r = p.add_run(_text(header.find(class_="cv-name")))
        r.bold, r.font.size = True, Pt(22)
        headline = _text(header.find(class_="cv-headline"))
        if headline:
            p = para(after=1, align=WD_ALIGN_PARAGRAPH.CENTER if centred else None)
            r = p.add_run(headline)
            r.font.size, r.italic = Pt(11.5), centred
            r.font.color.rgb = muted if template != "modern" else heading_color
        contact = header.find(class_="cv-contact")
        if contact:
            parts = [_text(s) for s in contact.find_all("span", recursive=False)]
            p = para(after=4, align=WD_ALIGN_PARAGRAPH.CENTER if centred else None)
            r = p.add_run(("  |  " if centred else "  ·  ").join(x for x in parts if x))
            r.font.size, r.font.color.rgb = Pt(9.5), muted

    for sec in soup.find_all("section", class_="cv-section"):
        if sec.has_attr("hidden"):
            continue
        h = para(before=8, after=3)
        r = h.add_run(_text(sec.find("h2")) if centred else _text(sec.find("h2")).upper())
        r.bold, r.font.size, r.font.color.rgb = True, Pt(10.5 if centred else 10), heading_color
        r.font.small_caps = centred
        _rule_below(h, "1B1B1B" if centred else "C9CFD8")

        for node in sec.find_all(["div", "p", "ul"], recursive=False):
            classes = node.get("class") or []
            if "cv-entry" in classes:
                two_sides([(_text(node.find(class_="cv-entry-title")), True)], _text(node.find(class_="cv-entry-dates")))
                sub = node.find(class_="cv-entry-sub")
                if sub and (_text(sub.find(class_="cv-entry-org")) or _text(sub.find(class_="cv-entry-loc"))):
                    two_sides([(_text(sub.find(class_="cv-entry-org")), False)], _text(sub.find(class_="cv-entry-loc")),
                              italic=True)
                desc = node.find(class_="cv-entry-desc")
                if desc:
                    para().add_run(_text(desc)).font.color.rgb = muted
                for li in node.find_all("li"):
                    b = doc.add_paragraph(style="List Bullet")
                    _tight(b, 0, 1)
                    b.add_run(_text(li))
                doc.paragraphs[-1].paragraph_format.space_after = Pt(4)
            elif "cv-items" in classes:
                for li in node.find_all("li"):
                    p = para(after=1)
                    strong = li.find("strong")
                    if strong:
                        p.add_run(_text(strong) + " ").bold = True
                        strong.extract()
                    p.add_run(_text(li))
            elif node.name in ("p", "div") and _text(node):
                para(after=2).add_run(_text(node))

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def letter_to_docx(body: str, *, template: str, accent: str, page_size: str) -> bytes:
    """A Word version of the cover letter page."""
    soup = BeautifulSoup(body, "html.parser")
    doc = Document()
    width_mm, height_mm = PAGE_MM["letter" if page_size == "letter" else "A4"]
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(width_mm), Mm(height_mm)
    section.left_margin = section.right_margin = Mm(22)
    section.top_margin = section.bottom_margin = Mm(20)
    normal = doc.styles["Normal"]
    normal.font.name = FONTS.get(template, "Arial")
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), normal.font.name)
    normal.font.size = Pt(11)
    muted = RGBColor(0x4F, 0x55, 0x61)

    header = soup.find(class_="cv-header")
    if header:
        p = doc.add_paragraph()
        _tight(p)
        r = p.add_run(_text(header.find(class_="cv-name")))
        r.bold, r.font.size = True, Pt(18)
        contact = header.find(class_="cv-contact")
        if contact:
            p = doc.add_paragraph()
            _tight(p, 0, 14)
            r = p.add_run("  ·  ".join(_text(s) for s in contact.find_all("span")))
            r.font.size, r.font.color.rgb = Pt(9.5), muted
            _rule_below(p, "C9CFD8")
        header.extract()
    for node in soup.find_all("p"):
        p = doc.add_paragraph()
        _tight(p, 0, 9)
        p.paragraph_format.line_spacing = 1.2
        lines = [line.strip() for line in node.get_text("\n").split("\n") if line.strip()]
        for i, line in enumerate(lines):
            run = p.add_run((" " if node.find("br") is None and i else "") + line)
            run.bold = bool(node.find("strong")) or "letter-sign" in (node.get("class") or [])
            if i < len(lines) - 1 and node.find("br") is not None:
                run.add_break()
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def docx_to_text(data: bytes) -> str:
    """The text of a Word document, including tables (many CV templates are laid out in tables)."""
    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:  # python-docx raises several error types for files that aren't .docx
        raise ExportError(f"Couldn't read that Word file: {e}") from e
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                if cell.text.strip() and cell.text not in cells:  # merged cells repeat their text
                    cells.append(cell.text)
            lines.append(" | ".join(cells))
    return "\n".join(line for line in lines if line.strip())


# ---- plain text (for pasting into application forms) --------------------

def html_to_text(body: str) -> str:
    """The CV as plain text, in reading order, for application forms that want text."""
    soup = BeautifulSoup(body, "html.parser")
    for auto in soup.select("[data-auto]"):
        auto.decompose()
    out: list[str] = []
    header = soup.find(class_="cv-header")
    if header:
        out.append(_text(header.find(class_="cv-name")))
        if _text(header.find(class_="cv-headline")):
            out.append(_text(header.find(class_="cv-headline")))
        contact = [_text(s) for s in header.select(".cv-contact > span")]
        if contact:
            out.append(" | ".join(c for c in contact if c))
    for sec in soup.find_all("section", class_="cv-section"):
        if sec.has_attr("hidden"):
            continue
        out += ["", _text(sec.find("h2")).upper()]
        for node in sec.find_all(["div", "p", "ul"], recursive=False):
            classes = node.get("class") or []
            if "cv-entry" in classes:
                title, dates = _text(node.find(class_="cv-entry-title")), _text(node.find(class_="cv-entry-dates"))
                org, loc = _text(node.find(class_="cv-entry-org")), _text(node.find(class_="cv-entry-loc"))
                out.append(" | ".join(x for x in (title, dates) if x))
                if org or loc:
                    out.append(", ".join(x for x in (org, loc) if x))
                if _text(node.find(class_="cv-entry-desc")):
                    out.append(_text(node.find(class_="cv-entry-desc")))
                out += [f"- {_text(li)}" for li in node.find_all("li")]
                out.append("")
            elif "cv-items" in classes:
                out += [_text(li) for li in node.find_all("li")]
            elif _text(node):
                out.append(_text(node))
    text = "\n".join(out)
    while "\n\n\n" in text:
        text = text.replace("\n\n\n", "\n\n")
    return text.strip() + "\n"
