import io

import pytest

from cv_maker.ai import AIError, ClaudeAI
from cv_maker.app import clean_html, create_app, select_ai
from cv_maker.ollama import OllamaAI
from cv_maker.schema import CVDocument, CVEntry, CVSection, Experience, IngestResult, Memory
from cv_maker.store import Store, ensure_ids


class FakeAI:
    """Stands in for Claude: appends each input as a new job and renders a simple CV."""

    def __init__(self):
        self.ingested = []
        self.targets = []

    def ingest(self, memory, text, attachments=None):
        self.ingested.append((text, [a.filename for a in attachments or []]))
        updated = memory.model_copy(deep=True)
        updated.profile.name = "Ada Lovelace"
        updated.experience.append(Experience(id="exp-new", role=text[:40] or "Imported", organization="Acme"))
        return IngestResult(memory=updated, changes=[f"Added {text[:20]}"], questions=["What was the impact?"])

    def improve_bullet(self, memory, bullet, mode, target="", instruction=""):
        self.improved = (bullet, mode, target, instruction)
        return [f"{mode}: {bullet}", f"Led {bullet.lower()}", bullet]

    def write_letter(self, memory, target, company, role, tone):
        from cv_maker.writing import CoverLetter
        self.letter_args = (target, company, role, tone)
        return CoverLetter(paragraphs=[f"I'd like to join {company or 'your team'} as {role or 'an engineer'}.",
                                       "At Acme I built X for 3 teams."])

    def build_cv(self, memory, target="", conventions=""):
        self.targets.append(target)
        self.conventions = conventions
        return CVDocument(
            name=memory.profile.name or "Nobody",
            sections=[CVSection(heading="Experience", entries=[
                CVEntry(title=e.role, subtitle=e.organization) for e in memory.experience
            ])],
            advice=["Add numbers"],
        )


    def truth_check(self, memory, cv_text):
        from cv_maker.assistant import TruthIssue, TruthReport
        self.checked = cv_text
        return TruthReport(issues=[TruthIssue(quote="Led 40 people", problem="The memory says you led 4.",
                                              suggestion="Led 4 people"), TruthIssue()])

    def strengthen_questions(self, memory, target="", count=6):
        from cv_maker.assistant import StrengthenQuestion, StrengthenQuestions
        return StrengthenQuestions(questions=[StrengthenQuestion(about="Acme", question="How many users?"),
                                              StrengthenQuestion(question=" ")])

    def interview_prep(self, memory, target, company, role):
        from cv_maker.assistant import InterviewPrep, InterviewQuestion
        self.prep_args = (target, company, role)
        return InterviewPrep(questions=[InterviewQuestion(question=f"Why {company}?", why="Always asked",
                                                          answer=["Your mission", "My Acme work"])],
                             ask_them=["What does success look like?"])

    def follow_up_email(self, memory, company, role, days, notes=""):
        from cv_maker.assistant import Email
        self.follow_up_args = (company, role, days, notes)
        return Email(subject=f"{role} application", body=f"Dear {company}, it's been {days} days.")

    def linkedin(self, memory, target=""):
        from cv_maker.assistant import LinkedInProfile, LinkedInRole
        return LinkedInProfile(headline="Data scientist", about="I like data.",
                               experience=[LinkedInRole(title="Engineer", company="Acme", description="Built X.")])

    def translate_cv(self, cv, language):
        self.translated_to = language
        out = cv.model_copy(deep=True)
        for section in out.sections:
            section.heading = "工作經驗" if section.heading == "Experience" else section.heading
        return out


@pytest.fixture
def ai():
    return FakeAI()


@pytest.fixture
def client(tmp_path, ai):
    return create_app(tmp_path, ai=ai).test_client()


def test_ingest_updates_memory_and_rebuilds_cv(client, ai):
    res = client.post("/api/ingest", data={"text": "Senior Engineer"})
    assert res.status_code == 200
    body = res.get_json()
    assert body["memory"]["experience"][0]["role"] == "Senior Engineer"
    assert body["changes"] == ["Added Senior Engineer"]
    assert body["questions"] == ["What was the impact?"]
    assert "Senior Engineer" in body["cv_html"]
    assert body["advice"] == ["Add numbers"]
    assert body["history"][0]["input"] == "Senior Engineer"


def test_memory_persists_between_app_instances(tmp_path, ai):
    create_app(tmp_path, ai=ai).test_client().post("/api/ingest", data={"text": "Engineer"})
    body = create_app(tmp_path, ai=ai).test_client().get("/api/state").get_json()
    assert body["memory"]["profile"]["name"] == "Ada Lovelace"
    assert len(body["history"]) == 1


def test_file_upload_is_passed_to_ai(client, ai):
    data = {"text": "", "files": (io.BytesIO(b"%PDF-1.4 fake"), "old_cv.pdf", "application/pdf")}
    res = client.post("/api/ingest", data=data, content_type="multipart/form-data")
    assert res.status_code == 200
    assert ai.ingested[-1][1] == ["old_cv.pdf"]


def test_unsupported_upload_is_rejected(client):
    data = {"files": (io.BytesIO(b"PK..."), "cv.pages", "application/octet-stream")}
    res = client.post("/api/ingest", data=data, content_type="multipart/form-data")
    assert res.status_code == 400 and "PDF, Word" in res.get_json()["error"]
    data = {"files": (io.BytesIO(b"not really word"), "cv.docx", "application/octet-stream")}
    res = client.post("/api/ingest", data=data, content_type="multipart/form-data")
    assert res.status_code == 400 and "Word file" in res.get_json()["error"]


def test_word_cv_is_read_as_text(client, ai):
    import docx
    doc = docx.Document()
    doc.add_paragraph("Ada Lovelace")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Analyst, Acme"
    table.rows[0].cells[1].text = "2021 – 2024"
    buf = io.BytesIO()
    doc.save(buf)
    seen = []
    original = ai.ingest
    ai.ingest = lambda memory, text, attachments=None: seen.extend(attachments or []) or original(memory, text, attachments)
    data = {"files": (io.BytesIO(buf.getvalue()), "My CV.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    assert client.post("/api/ingest", data=data, content_type="multipart/form-data").status_code == 200
    assert seen[0].media_type == "text/plain"
    text = seen[0].data.decode()
    assert "Ada Lovelace" in text and "Analyst, Acme | 2021 – 2024" in text


def test_empty_input_is_rejected(client):
    assert client.post("/api/ingest", data={"text": "  "}).status_code == 400


def test_auto_rebuild_can_be_turned_off(client, ai):
    client.post("/api/settings", json={"auto_rebuild": False})
    body = client.post("/api/ingest", data={"text": "Engineer"}).get_json()
    assert body["cv_html"] == ""
    assert ai.targets == []


def test_manual_edits_are_not_overwritten_by_auto_rebuild(client):
    client.post("/api/ingest", data={"text": "Engineer"})
    client.post("/api/cv", json={"html": "<h1>My own wording</h1>"})
    body = client.post("/api/ingest", data={"text": "Manager"}).get_json()
    assert body["cv_html"] == "<h1>My own wording</h1>"
    assert "manual edits" in body["notice"]
    assert body["cv_meta"]["edited"] is True


def test_learning_from_edits_updates_memory_and_clears_flag(client, ai):
    client.post("/api/ingest", data={"text": "Engineer"})
    client.post("/api/cv", json={"html": "<p>edited</p>"})
    body = client.post("/api/cv/learn", json={"text": "Engineer at Acme, 2020-2024"}).get_json()
    assert body["cv_meta"]["edited"] is False
    assert "edited their CV by hand" in ai.ingested[-1][0]


def test_build_uses_and_saves_target(client, ai):
    client.post("/api/ingest", data={"text": "Engineer"})
    body = client.post("/api/build", json={"target": "Staff engineer at a fintech"}).get_json()
    assert ai.targets[-1] == "Staff engineer at a fintech"
    assert body["settings"]["target"] == "Staff engineer at a fintech"
    assert body["cv_meta"]["target"] == "Staff engineer at a fintech"


def test_undo_restores_previous_memory(client):
    client.post("/api/ingest", data={"text": "First"})
    client.post("/api/ingest", data={"text": "Second"})
    body = client.post("/api/undo").get_json()
    assert [e["role"] for e in body["memory"]["experience"]] == ["First"]
    assert "First" in body["cv_html"] and "Second" not in body["cv_html"]
    client.post("/api/undo")
    assert client.post("/api/undo").status_code == 400


def test_manual_memory_edit_is_validated(client):
    assert client.put("/api/memory", json={"experience": "not a list"}).status_code == 400
    body = client.put("/api/memory", json={"profile": {"name": "Grace"}}).get_json()
    assert body["memory"]["profile"]["name"] == "Grace"


def test_without_ai_builds_plain_cv_and_refuses_ingest(tmp_path):
    client = create_app(tmp_path / "data", ai=None).test_client()
    state = client.get("/api/state").get_json()
    assert state["ai_enabled"] is False
    assert state["ai_status"]["ready"] is False
    assert client.post("/api/ingest", data={"text": "hi"}).status_code == 400
    client.put("/api/memory", json={"profile": {"name": "Grace"}, "skills": [{"category": "Languages", "skills": ["COBOL"]}]})
    html = client.get("/api/state").get_json()["cv_html"]
    assert "Grace" in html and "COBOL" in html


def test_auto_backend_uses_free_local_model_without_api_key(tmp_path, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.delenv("ANTHROPIC_AUTH_TOKEN", raising=False)
    monkeypatch.setenv("HOME", str(tmp_path))
    assert isinstance(select_ai("auto"), OllamaAI)
    assert select_ai("none") is None
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-test")
    assert isinstance(select_ai("auto"), ClaudeAI)
    assert isinstance(select_ai("api"), ClaudeAI)
    assert isinstance(select_ai("ollama"), OllamaAI)


def test_failed_auto_rebuild_still_saves_memory(tmp_path, ai):
    def broken(*args, **kwargs):
        raise AIError("Can't reach Ollama.")

    ai.build_cv = broken
    client = create_app(tmp_path, ai=ai).test_client()
    body = client.post("/api/ingest", data={"text": "Engineer"}).get_json()
    assert body["memory"]["experience"][0]["role"] == "Engineer"
    assert "couldn't be rebuilt" in body["notice"]


def test_standalone_cv_page_escapes_name(client, ai):
    client.put("/api/memory", json={"profile": {"name": "<b>x</b>"}})
    page = client.get("/cv.html").get_data(as_text=True)
    assert "<title>&lt;b&gt;x&lt;/b&gt; – CV</title>" in page


def test_clean_html_strips_scripts_and_handlers():
    dirty = '<p onclick="x()">Hi</p><script>alert(1)</script><a href="javascript:bad()">l</a>'
    cleaned = clean_html(dirty)
    assert "script" not in cleaned and "onclick" not in cleaned and "javascript" not in cleaned
    assert "<p>Hi</p>" in cleaned


def test_ensure_ids_fills_and_dedupes():
    memory = Memory(experience=[Experience(id="Acme Job"), Experience(id="acme-job"), Experience(id="")])
    ids = [e.id for e in ensure_ids(memory).experience]
    assert ids == ["acme-job", "acme-job-2", "exp-3"]


def test_store_history_is_newest_first(tmp_path):
    store = Store(tmp_path)
    store.save_memory(Memory(), source="input", input_text="one")
    store.save_memory(Memory(), source="input", input_text="two")
    assert [h["input"] for h in store.history()] == ["two", "one"]


def test_section_layout_is_saved_and_survives_rebuilds(client, ai):
    client.post("/api/ingest", data={"text": "Engineer"})
    client.put("/api/memory?rebuild=0", json={
        "profile": {"name": "Ada"}, "experience": [{"id": "x", "role": "Engineer"}],
        "skills": [{"category": "Languages", "skills": ["Python"]}],
    })
    ai.build_cv = lambda memory, target="", conventions="": CVDocument(name="Ada", sections=[
        CVSection(heading="Experience", items=["job"]), CVSection(heading="Skills", items=["Python"]),
    ])
    client.post("/api/build", json={})
    body = client.post("/api/cv/layout", json={"order": ["skills", "experience"], "hidden": []}).get_json()
    assert [s["key"] for s in body["sections"]] == ["skills", "experience"]
    assert body["cv_html"].index('data-section="skills"') < body["cv_html"].index('data-section="experience"')
    assert body["cv_meta"]["edited"] is False

    body = client.post("/api/build", json={}).get_json()  # a rebuild keeps the chosen order
    assert [s["key"] for s in body["sections"]] == ["skills", "experience"]
    body = client.post("/api/cv/layout", json={"order": ["skills", "experience"], "hidden": ["skills"]}).get_json()
    assert body["sections"][0]["hidden"] is True and 'data-section="skills" hidden' in body["cv_html"]


def test_layout_change_keeps_hand_edits(client):
    client.post("/api/ingest", data={"text": "Engineer"})
    client.post("/api/cv", json={"html": "<section data-section='experience'><p>my words</p></section>"})
    body = client.post("/api/cv/layout", json={
        "order": ["experience"], "hidden": [], "html": "<section data-section='experience'><p>my words, moved</p></section>",
    }).get_json()
    assert "my words, moved" in body["cv_html"] and body["cv_meta"]["edited"] is True


def test_memory_tab_edits_do_not_trigger_a_rebuild(client, ai):
    body = client.put("/api/memory?rebuild=0", json={"profile": {"name": "Grace"}}).get_json()
    assert ai.targets == [] and "Rebuild CV" in body["notice"]
    assert body["memory"]["profile"]["name"] == "Grace"


def test_design_settings_are_validated_and_used_on_the_standalone_page(client):
    assert client.post("/api/settings", json={"template": "fancy"}).status_code == 400
    assert client.post("/api/settings", json={"accent": "red; background: url(x)"}).status_code == 400
    body = client.post("/api/settings", json={"template": "modern", "accent": "#0f6e6e", "fit_one_page": 1}).get_json()
    assert body["settings"]["template"] == "modern" and body["settings"]["fit_one_page"] is True
    client.post("/api/ingest", data={"text": "Engineer"})
    page = client.get("/cv.html").get_data(as_text=True)
    assert "class='cv t-modern'" in page and "--cv-accent:#0f6e6e" in page


def test_improve_a_bullet(client, ai):
    client.post("/api/versions", json={"company": "HSBC", "target": "Python role"})
    res = client.post("/api/improve", json={"text": "Built a dashboard", "mode": "job"})
    assert res.status_code == 200 and res.get_json()["suggestions"][0] == "job: Built a dashboard"
    assert ai.improved == ("Built a dashboard", "job", "Python role", "")
    assert client.post("/api/improve", json={"text": "", "mode": "stronger"}).status_code == 400
    assert client.post("/api/improve", json={"text": "x", "mode": "rhyme"}).status_code == 400


def test_improve_needs_an_ai(tmp_path):
    client = create_app(tmp_path, ai=None).test_client()
    assert client.post("/api/improve", json={"text": "Built X", "mode": "stronger"}).status_code == 400


def test_suggestions_drop_duplicates_and_bullet_marks():
    from cv_maker.writing import BulletSuggestions, clean_suggestions
    out = clean_suggestions(BulletSuggestions(suggestions=["• Led X", "led x", "Built X", " - Shipped Y", "Z", "W"]), "Built X")
    assert out == ["Led X", "Shipped Y", "Z"]


def test_other_websites_cannot_change_anything(tmp_path, ai):
    from flask.testing import FlaskClient
    app = create_app(tmp_path, ai=ai)
    outsider = FlaskClient(app)  # a plain browser request, e.g. a form on another website
    res = outsider.post("/api/ingest", data={"text": "I am the CEO of Google"})
    assert res.status_code == 403 and ai.ingested == []
    assert outsider.post("/api/build", json={}).status_code == 403
    assert outsider.get("/api/state").status_code == 200  # reading is fine; browsers keep it from other sites


def test_requests_for_other_host_names_are_refused(tmp_path, ai):
    client = create_app(tmp_path, ai=ai).test_client()
    assert client.get("/api/state", headers={"Host": "evil.example.com"}).status_code == 403
    assert client.get("/api/state", headers={"Host": "127.0.0.1:5000"}).status_code == 200


def test_undo_my_edits_restores_the_generated_cv(client):
    original = client.post("/api/ingest", data={"text": "Engineer"}).get_json()["cv_html"]
    client.post("/api/cv", json={"html": "<p>oops</p>"})
    body = client.post("/api/cv/reset").get_json()
    assert body["cv_html"] == original and body["cv_meta"]["edited"] is False


def test_backup_contains_everything_but_the_api_key(tmp_path, client):
    import io
    import zipfile
    client.post("/api/ingest", data={"text": "Engineer"})
    client.post("/api/versions", json={"company": "Acme"})
    (tmp_path / "secrets.json").write_text('{"key": "sk-secret"}')
    res = client.get("/api/export/backup.zip")
    names = zipfile.ZipFile(io.BytesIO(res.data)).namelist()
    assert "data/memory.json" in names and "data/cv.html" in names
    assert any(n.startswith("data/versions/") and n.endswith("info.json") for n in names)
    assert not any("secrets" in n for n in names) and b"sk-secret" not in res.data
